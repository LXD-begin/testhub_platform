import json
from dataclasses import dataclass

from openpyxl import load_workbook


CASE_COLUMNS = [
    'case_id',
    'module',
    'title',
    'enabled',
    'method',
    'path',
    'headers',
    'params',
    'body',
    'extract',
    'expected_status',
    'expected_code',
    'expected_message',
    'expected_fields',
    'depends_on',
    'description',
]


@dataclass
class ApiCase:
    case_id: str
    module: str
    title: str
    enabled: bool
    method: str
    path: str
    headers: dict
    params: dict
    body: object
    extract: dict
    expected_status: int
    expected_code: int | None
    expected_message: str
    expected_fields: dict
    depends_on: str
    description: str


def parse_json_cell(value, default):
    if value in (None, ''):
        return default
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(str(value))
    except json.JSONDecodeError as exc:
        raise ValueError(f'Excel JSON 单元格格式错误: {value}') from exc


def parse_bool(value):
    return str(value).strip().lower() in {'1', 'true', 'yes', 'y', '是', '启用'}


def load_cases(file_path):
    workbook = load_workbook(file_path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    missing = [column for column in CASE_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f'Excel 缺少列: {", ".join(missing)}')

    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        raw = dict(zip(headers, row))
        case_id = raw.get('case_id')
        if not case_id:
            continue
        rows.append(
            ApiCase(
                case_id=str(case_id),
                module=str(raw.get('module') or ''),
                title=str(raw.get('title') or case_id),
                enabled=parse_bool(raw.get('enabled', True)),
                method=str(raw.get('method') or 'GET').upper(),
                path=str(raw.get('path') or ''),
                headers=parse_json_cell(raw.get('headers'), {}),
                params=parse_json_cell(raw.get('params'), {}),
                body=parse_json_cell(raw.get('body'), None),
                extract=parse_json_cell(raw.get('extract'), {}),
                expected_status=int(raw.get('expected_status') or 200),
                expected_code=int(raw['expected_code']) if raw.get('expected_code') not in (None, '') else None,
                expected_message=str(raw.get('expected_message') or ''),
                expected_fields=parse_json_cell(raw.get('expected_fields'), {}),
                depends_on=str(raw.get('depends_on') or ''),
                description=str(raw.get('description') or ''),
            )
        )
    return [case for case in rows if case.enabled]
