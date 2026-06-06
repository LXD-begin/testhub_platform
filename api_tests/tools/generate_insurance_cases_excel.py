import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from api_tests.common.excel_reader import CASE_COLUMNS  # noqa: E402
from api_tests.config import CASE_FILE  # noqa: E402


def dumps(value):
    if value in (None, ''):
        return ''
    return json.dumps(value, ensure_ascii=False)


def row(
    case_id,
    module,
    title,
    method,
    path,
    *,
    headers=None,
    params=None,
    body=None,
    extract=None,
    expected_status=200,
    expected_code=200,
    expected_message='成功',
    expected_fields=None,
    depends_on='',
    description='',
    enabled=True,
):
    return {
        'case_id': case_id,
        'module': module,
        'title': title,
        'enabled': '是' if enabled else '否',
        'method': method,
        'path': path,
        'headers': dumps(headers or {}),
        'params': dumps(params or {}),
        'body': dumps(body),
        'extract': dumps(extract or {}),
        'expected_status': expected_status,
        'expected_code': expected_code,
        'expected_message': expected_message,
        'expected_fields': dumps(expected_fields or {}),
        'depends_on': depends_on,
        'description': description,
    }


def premium_trial_body(**overrides):
    body = {
        'product_code': 'PA_C_ACCIDENT',
        'plan_code': 'STANDARD',
        'effective_date': '2026-06-06',
        'insurance_period_months': 12,
        'applicant': {
            'name': '张三',
            'id_type': 'IDENTITY_CARD',
            'id_no': '110101199001011234',
            'mobile': '13800138000',
            'email': 'zhangsan@example.com',
        },
        'insured': {
            'name': '张三',
            'id_type': 'IDENTITY_CARD',
            'id_no': '110101199001011234',
            'mobile': '13800138000',
            'email': 'zhangsan@example.com',
        },
        'occupation_code': '010101',
        'occupation_category': 2,
        'has_social_security': True,
        'channel_code': 'C_APP',
    }
    body.update(overrides)
    return body


def build_cases():
    auth_header = {'Authorization': 'Bearer ${token}'}
    return [
        row(
            'AUTH_001',
            '鉴权',
            '获取随机 token',
            'POST',
            '/api/insurance/auth/token/',
            extract={'token': 'data.access_token'},
            expected_fields={'data.token_type': 'Bearer'},
            description='无需登录，每次返回随机 token。',
        ),
        row(
            'PRODUCT_001',
            '产品配置',
            '查询产品列表',
            'GET',
            '/api/insurance/products/',
            headers=auth_header,
            expected_fields={'data.0.product_code': 'PA_C_ACCIDENT'},
            depends_on='AUTH_001',
        ),
        row(
            'QUOTE_001',
            '保费试算',
            '保费试算成功',
            'POST',
            '/api/insurance/premium-trials/',
            headers=auth_header,
            body=premium_trial_body(),
            extract={'quote_no': 'data.quote_no', 'payable_premium': 'data.premium_detail.payable_premium'},
            expected_status=201,
            expected_fields={'data.status': 'QUOTED', 'data.plan_code': 'STANDARD'},
            depends_on='PRODUCT_001',
        ),
        row(
            'QUOTE_002',
            '保费试算',
            '缺少计划代码',
            'POST',
            '/api/insurance/premium-trials/',
            headers=auth_header,
            body={key: value for key, value in premium_trial_body().items() if key != 'plan_code'},
            expected_status=400,
            expected_code=999,
            expected_message='参数校验失败',
            description='单接口异常：必填字段缺失。',
            depends_on='AUTH_001',
        ),
        row(
            'QUOTE_003',
            '保费试算',
            '计划代码不存在',
            'POST',
            '/api/insurance/premium-trials/',
            headers=auth_header,
            body=premium_trial_body(plan_code='VIP'),
            expected_status=400,
            expected_code=999,
            expected_message='产品计划不存在或未启用',
            expected_fields={'data.error_type': 'BUSINESS_ERROR'},
            description='单接口异常：业务参数非法。',
            depends_on='AUTH_001',
        ),
        row(
            'APP_001',
            '投保单',
            '创建投保单成功',
            'POST',
            '/api/insurance/applications/',
            headers=auth_header,
            body={
                'quote_no': '${quote_no}',
                'disclosures': {
                    'truth_declaration_confirmed': True,
                    'health_answers': {
                        'has_major_disease': False,
                        'has_disability': False,
                        'has_recent_claim': False,
                    },
                },
                'beneficiary_type': 'LEGAL',
                'beneficiaries': [],
                'consents': {
                    'terms_confirmed': True,
                    'exclusions_confirmed': True,
                    'electronic_policy_confirmed': True,
                },
                'channel_code': 'C_APP',
            },
            extract={'application_no': 'data.application_no'},
            expected_status=201,
            expected_fields={'data.status': 'CREATED'},
            depends_on='QUOTE_001',
        ),
        row(
            'PAY_001',
            '支付',
            '未核保不能创建支付订单',
            'POST',
            '/api/insurance/payment-orders/',
            headers=auth_header,
            body={'application_no': '${application_no}', 'pay_channel': 'MOCK'},
            expected_status=400,
            expected_code=999,
            expected_message='投保单未核保通过，不能创建支付订单',
            depends_on='APP_001',
            description='多接口异常：投保单未核保直接支付。',
        ),
        row(
            'UW_001',
            '核保',
            '投保单提交核保成功',
            'POST',
            '/api/insurance/applications/${application_no}/underwriting/',
            headers=auth_header,
            body={},
            extract={'underwriting_no': 'data.uw_no'},
            expected_status=201,
            expected_fields={'data.decision': 'APPROVED'},
            depends_on='APP_001',
        ),
        row(
            'PAY_002',
            '支付',
            '创建支付订单成功',
            'POST',
            '/api/insurance/payment-orders/',
            headers=auth_header,
            body={'application_no': '${application_no}', 'pay_channel': 'MOCK'},
            extract={'pay_order_no': 'data.pay_order_no', 'pay_amount': 'data.amount'},
            expected_status=201,
            expected_fields={'data.status': 'CREATED'},
            depends_on='UW_001',
        ),
        row(
            'PAY_003',
            '支付',
            '支付回调成功',
            'POST',
            '/api/insurance/payment-orders/${pay_order_no}/confirm/',
            headers=auth_header,
            body={
                'pay_status': 'SUCCESS',
                'paid_amount': '${pay_amount}',
                'external_trade_no': 'MOCK202606060001',
                'pay_channel': 'MOCK',
            },
            expected_fields={'data.status': 'SUCCESS'},
            depends_on='PAY_002',
        ),
        row(
            'POLICY_001',
            '承保出单',
            '按投保单承保出单成功',
            'POST',
            '/api/insurance/application-policies/',
            headers=auth_header,
            body={
                'application_no': '${application_no}',
                'delivery': {'email': 'zhangsan@example.com', 'sms_mobile': '13800138000'},
            },
            extract={'policy_no': 'data.policy_no'},
            expected_status=201,
            expected_fields={'data.status': 'ISSUED'},
            depends_on='PAY_003',
        ),
        row(
            'POLICY_002',
            '保单查询',
            '查询保单详情成功',
            'GET',
            '/api/insurance/policies/${policy_no}/',
            headers=auth_header,
            expected_fields={'data.policy_no': '${policy_no}', 'data.status': 'ISSUED'},
            depends_on='POLICY_001',
        ),
    ]


def generate(file_path=CASE_FILE):
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'insurance_api_cases'
    sheet.append(CASE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
        cell.alignment = Alignment(horizontal='center')

    for case in build_cases():
        sheet.append([case[column] for column in CASE_COLUMNS])

    for column_cells in sheet.columns:
        column = column_cells[0].column_letter
        sheet.column_dimensions[column].width = 24
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    workbook.save(file_path)
    return file_path


if __name__ == '__main__':
    path = generate()
    print(f'已生成接口用例 Excel: {path}')
